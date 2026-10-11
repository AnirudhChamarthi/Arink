"""Step 11 security — session isolation, report paths, Excel/chart bounds."""

import os

import numpy as np
import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from main import app
from reports.excel_writer import validate_report_path, write_nlq_report
from reports.report_charts import build_nlq_charts
from schemas.contract import AssetClass, EnrichedPortfolio, EnrichedPosition

_PNG_HEADER = b"\x89PNG\r\n\x1a\n"


def _valid_chart_png() -> bytes:
    pos = EnrichedPosition(
        ticker="AAPL",
        name="Apple",
        asset_class=AssetClass.EQUITY,
        quantity=10.0,
        cost_price=100.0,
        currency="USD",
        sector="Technology",
        current_price=100.0,
        returns_252d=np.random.default_rng(0).normal(0, 0.01, 252).tolist(),
        implied_vol=0.25,
        risk_free_rate=0.04,
        esg_score=70.0,
        carbon_intensity=40.0,
    )
    portfolio = EnrichedPortfolio(
        portfolio_name="Sec",
        base_currency="USD",
        positions=[pos],
    )
    charts = build_nlq_charts(
        models_used=["var:parametric"],
        detail_rows=[{"ticker": "AAPL", "pct_of_total": 1.0}],
        portfolio=portfolio,
    )
    assert charts
    return charts[0][1]


@pytest.mark.asyncio
async def test_risk_routes_reject_unknown_session():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        bogus = "00000000-0000-0000-0000-000000000000"
        for path, payload in [
            ("/risk/var", {"portfolio_session_id": bogus, "method": "parametric", "confidence": 0.99, "horizon_days": 1}),
            ("/risk/greeks", {"portfolio_session_id": bogus}),
            ("/risk/full-report", {"portfolio_session_id": bogus}),
        ]:
            r = await client.post(path, json=payload)
            assert r.status_code == 422, path
            assert "portfolio_session_id" in r.text.lower()


def test_validate_report_path_rejects_non_xlsx(tmp_path, monkeypatch):
    monkeypatch.setattr("reports.excel_writer._REPORTS_DIR", tmp_path)
    bad = tmp_path / "report.txt"
    bad.write_text("not excel")
    with pytest.raises(ValueError, match="xlsx"):
        validate_report_path(str(bad))


def test_validate_report_path_rejects_parent_traversal(tmp_path, monkeypatch):
    reports = tmp_path / "reports"
    reports.mkdir()
    monkeypatch.setattr("reports.excel_writer._REPORTS_DIR", reports)
    outside = tmp_path / "outside.xlsx"
    outside.write_bytes(b"")
    traversal = reports / ".." / "outside.xlsx"
    with pytest.raises(ValueError, match="outside"):
        validate_report_path(str(traversal.resolve()))


def test_nlq_report_caps_detail_rows(tmp_path, monkeypatch):
    monkeypatch.setattr("reports.excel_writer._REPORTS_DIR", tmp_path)
    rows = [{"ticker": f"T{i:04d}", "weight": i} for i in range(1500)]
    path = write_nlq_report(
        slug="cap",
        topline_number="1%",
        topline_label="VaR",
        formula_breakdown=[],
        explanation="test",
        models_used=["var:parametric"],
        detail_rows=rows,
        query="cap test",
    )
    wb = load_workbook(path)
    detail = wb["Position Detail"]
    assert detail.max_row == 1001
    os.remove(path)


def test_nlq_report_caps_and_filters_charts(tmp_path, monkeypatch):
    monkeypatch.setattr("reports.excel_writer._REPORTS_DIR", tmp_path)
    valid = _valid_chart_png()
    oversized = _PNG_HEADER + (b"\x00" * (2 * 1024 * 1024 + 1))
    charts = (
        [("bad", b"NOTPNG")] + [(f"chart_{i}", valid) for i in range(12)] + [("huge", oversized)]
    )
    path = write_nlq_report(
        slug="charts",
        topline_number="1%",
        topline_label="VaR",
        formula_breakdown=[],
        explanation="test",
        models_used=["var:parametric"],
        detail_rows=[],
        query="charts",
        chart_images=charts,
    )
    wb = load_workbook(path)
    assert len(wb["Charts"]._images) <= 8
    os.remove(path)
