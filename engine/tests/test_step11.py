"""Step 11 — model unit tests, API smoke, Excel workbook structure."""

import os
from pathlib import Path

import numpy as np
import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from main import app
from models.garch import run_garch
from models.greeks import run_greeks
from models.stress import run_stress
from reports.excel_writer import write_nlq_report
from reports.report_charts import build_nlq_charts
from schemas.contract import AssetClass, EnrichedPortfolio, EnrichedPosition, StressScenario


def _position(ticker: str, sector: str = "Technology") -> EnrichedPosition:
    rng = np.random.default_rng(42)
    return EnrichedPosition(
        ticker=ticker,
        name=f"{ticker} Inc",
        asset_class=AssetClass.EQUITY,
        quantity=10.0,
        cost_price=90.0,
        currency="USD",
        sector=sector,
        current_price=100.0,
        returns_252d=rng.normal(0, 0.01, 252).tolist(),
        implied_vol=0.25,
        risk_free_rate=0.04,
        esg_score=70.0,
        carbon_intensity=40.0,
        beta=1.0,
    )


def _portfolio() -> EnrichedPortfolio:
    return EnrichedPortfolio(
        portfolio_name="Step11",
        base_currency="USD",
        positions=[_position("AAPL"), _position("MSFT")],
    )


def test_greeks_deterministic_with_fixed_seed():
    p = _portfolio()
    g1 = run_greeks(p)
    g2 = run_greeks(p)
    assert g1.portfolio_delta == g2.portfolio_delta
    assert len(g1.positions) == 2


def test_stress_and_garch_models():
    p = _portfolio()
    stress = run_stress(p, StressScenario.CREDIT_2008)
    assert stress.portfolio_pnl < 0

    garch = run_garch(p, "AAPL")
    assert garch.ticker == "AAPL"
    assert garch.vol_forecast_10d > 0


def test_nlq_workbook_structure_snapshot(tmp_path, monkeypatch):
    monkeypatch.setattr("reports.excel_writer._REPORTS_DIR", tmp_path)
    charts = build_nlq_charts(
        models_used=["var:parametric"],
        detail_rows=[{"ticker": "AAPL", "pct_of_total": 0.6}],
        portfolio=_portfolio(),
    )
    path = write_nlq_report(
        slug="step11",
        topline_number="1.00%",
        topline_label="VaR",
        formula_breakdown=["VaR = z * sigma"],
        explanation="test",
        models_used=["var:parametric"],
        detail_rows=[{"ticker": "AAPL", "pct_of_total": 0.6}],
        query="What is VaR?",
        chart_images=charts,
        model_parameters={"confidence": 0.99},
    )
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Position Detail", "Model Parameters", "Charts"]
    assert len(wb["Charts"]._images) >= 1
    os.remove(path)


@pytest.mark.asyncio
async def test_api_smoke_all_risk_routes():
    body = {
        "portfolio_name": "Step11 API",
        "base_currency": "USD",
        "positions": [
            {
                "ticker": "AAPL",
                "asset_class": "Equity",
                "quantity": 10,
                "cost_price": 100,
                "currency": "USD",
            },
        ],
    }
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        health = await client.get("/health")
        assert health.status_code == 200

        enrich = await client.post("/enrich?response_detail=display", json=body)
        assert enrich.status_code == 200
        session_id = enrich.json()["portfolio_session_id"]

        for path, payload in [
            ("/risk/var", {"portfolio_session_id": session_id, "method": "parametric", "confidence": 0.99, "horizon_days": 1}),
            ("/risk/greeks", {"portfolio_session_id": session_id}),
            (
                "/risk/stress",
                {"portfolio_session_id": session_id, "scenario": "2008_credit"},
            ),
            (
                "/risk/sustainability",
                {"portfolio_session_id": session_id, "exclusion_filters": ["ticker:PLTR"]},
            ),
            ("/risk/garch", {"portfolio_session_id": session_id, "ticker": "AAPL"}),
        ]:
            r = await client.post(path, json=payload)
            assert r.status_code == 200, f"{path} failed: {r.text}"
            if path == "/risk/stress":
                report_path = r.json()["excel_report_path"]
                assert Path(report_path).is_file()
                os.remove(report_path)

        full = await client.post("/risk/full-report", json={"portfolio_session_id": session_id})
        assert full.status_code == 200
        report_path = full.json()["excel_report_path"]
        assert Path(report_path).is_file()
        os.remove(report_path)
