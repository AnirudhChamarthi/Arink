"""Step 9 — Excel report charts and full-risk workbook."""

import os
from pathlib import Path

import numpy as np
import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from main import app
from reports.excel_writer import write_nlq_report
from reports.report_charts import build_nlq_charts
from schemas.contract import AssetClass, EnrichedPortfolio, EnrichedPosition


def _position(ticker: str) -> EnrichedPosition:
    return EnrichedPosition(
        ticker=ticker,
        asset_class=AssetClass.EQUITY,
        quantity=10.0,
        cost_price=100.0,
        currency="USD",
        sector="Technology",
        current_price=100.0,
        returns_252d=np.random.default_rng(hash(ticker) % 2**32).normal(0, 0.012, 252).tolist(),
        implied_vol=0.25,
        risk_free_rate=0.04,
        esg_score=75.0,
        carbon_intensity=40.0,
    )


def _portfolio() -> EnrichedPortfolio:
    return EnrichedPortfolio(
        portfolio_name="Chart Test",
        base_currency="USD",
        positions=[_position("AAPL"), _position("MSFT")],
    )


def test_build_nlq_var_charts_returns_png_bytes():
    portfolio = _portfolio()
    detail = [
        {"ticker": "AAPL", "var_contribution": 0.01, "pct_of_total": 0.6},
        {"ticker": "MSFT", "var_contribution": 0.005, "pct_of_total": 0.4},
    ]
    charts = build_nlq_charts(
        models_used=["var:parametric"],
        detail_rows=detail,
        portfolio=portfolio,
    )
    assert len(charts) >= 2
    for title, png in charts:
        assert title
        assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_write_nlq_report_includes_charts_sheet(tmp_path, monkeypatch):
    monkeypatch.setattr("reports.excel_writer._REPORTS_DIR", tmp_path)
    portfolio = _portfolio()
    charts = build_nlq_charts(
        models_used=["var:parametric"],
        detail_rows=[{"ticker": "AAPL", "pct_of_total": 0.5}],
        portfolio=portfolio,
    )
    path = write_nlq_report(
        slug="test",
        topline_number="2.50%",
        topline_label="VaR",
        formula_breakdown=["line 1"],
        explanation="test",
        models_used=["var:parametric"],
        detail_rows=[{"ticker": "AAPL", "pct_of_total": 0.5}],
        query="VaR?",
        chart_images=charts,
        model_parameters={"confidence": 0.99},
    )
    wb = load_workbook(path)
    assert "Charts" in wb.sheetnames
    assert "Model Parameters" in wb.sheetnames
    assert len(wb["Charts"]._images) >= 1
    os.remove(path)


@pytest.mark.asyncio
async def test_nlq_report_has_charts_sheet():
    body = {
        "portfolio_name": "NLQ Charts",
        "base_currency": "USD",
        "positions": [
            {
                "ticker": "AAPL",
                "asset_class": "Equity",
                "quantity": 10,
                "cost_price": 100,
                "currency": "USD",
            },
            {
                "ticker": "MSFT",
                "asset_class": "Equity",
                "quantity": 5,
                "cost_price": 200,
                "currency": "USD",
            },
        ],
    }
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        enrich = await client.post("/enrich?response_detail=display", json=body)
        session_id = enrich.json()["portfolio_session_id"]
        nlq = await client.post(
            "/nlq",
            json={
                "query": "What is my 99% VaR?",
                "portfolio_session_id": session_id,
                "limits": {},
            },
        )
    assert nlq.status_code == 200
    path = nlq.json()["excel_report_path"]
    wb = load_workbook(path)
    assert "Charts" in wb.sheetnames
    assert len(wb["Charts"]._images) >= 1
    os.remove(path)


@pytest.mark.asyncio
async def test_full_report_endpoint():
    body = {
        "portfolio_name": "Full Report",
        "base_currency": "USD",
        "positions": [
            {
                "ticker": "AAPL",
                "asset_class": "Equity",
                "quantity": 100,
                "cost_price": 150,
                "currency": "USD",
            },
        ],
    }
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        enrich = await client.post("/enrich?response_detail=display", json=body)
        session_id = enrich.json()["portfolio_session_id"]
        report = await client.post(
            "/risk/full-report",
            json={"portfolio_session_id": session_id},
        )
    assert report.status_code == 200
    data = report.json()
    assert data["portfolio_name"] == "Full Report"
    assert "VaR" in data["summary_text"]
    path = data["excel_report_path"]
    assert Path(path).is_file()
    wb = load_workbook(path)
    assert "Charts" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert len(wb["Charts"]._images) >= 2
    os.remove(path)
