"""Generate Excel NLQ and full-risk reports."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font

from reports.excel_safety import excel_safe_cell

_ENGINE_ROOT = Path(__file__).resolve().parents[1]
_REPORTS_DIR = _ENGINE_ROOT / "output" / "reports"
_BOLD = Font(bold=True, size=14)
_BOLD_HEADER = Font(bold=True)
_MAX_DETAIL_ROWS = 1000
_MAX_CHARTS = 8
_MAX_CHART_PNG_BYTES = 2 * 1024 * 1024
_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def write_nlq_report(
    *,
    slug: str,
    topline_number: str,
    topline_label: str,
    formula_breakdown: list[str],
    explanation: str,
    models_used: list[str],
    detail_rows: list[dict[str, Any]],
    query: str,
    chart_images: list[tuple[str, bytes]] | None = None,
    model_parameters: dict[str, Any] | None = None,
) -> str:
    _REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = _REPORTS_DIR / f"riskbot_{slug}_{ts}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = topline_number
    ws["A1"].font = Font(bold=True, size=24)
    ws["A2"] = topline_label
    ws["A2"].font = _BOLD
    ws["A4"] = "Query"
    ws["B4"] = excel_safe_cell(query)
    ws["A5"] = "Models"
    ws["B5"] = ", ".join(models_used)
    ws["A7"] = "Formula breakdown"
    ws["A7"].font = _BOLD
    row = 8
    for line in formula_breakdown:
        ws.cell(row=row, column=1, value=excel_safe_cell(line))
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Explanation").font = _BOLD
    row += 1
    ws.cell(row=row, column=1, value=excel_safe_cell(explanation))
    ws.column_dimensions["A"].width = 80
    ws.column_dimensions["B"].width = 40

    if detail_rows:
        _write_detail_sheet(wb, detail_rows)

    _write_model_parameters_sheet(
        wb,
        model_parameters or {},
        path,
        extra={"query": query, "models": ", ".join(models_used)},
    )

    if chart_images:
        _write_charts_sheet(wb, chart_images)

    wb.save(path)
    return str(path.resolve())


def write_full_report(
    *,
    portfolio,
    var_result,
    greeks_result,
    sustainability_result,
    stress_result,
    chart_images: list[tuple[str, bytes]],
    model_parameters: dict[str, Any],
) -> str:
    _REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    slug = _slugify(portfolio.portfolio_name)
    path = _REPORTS_DIR / f"riskbot_full_{slug}_{ts}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = excel_safe_cell(portfolio.portfolio_name)
    ws["A1"].font = Font(bold=True, size=20)
    row = 3
    metrics = [
        ("VaR", f"{var_result.var * 100:.2f}%", f"{var_result.method.value}, {var_result.confidence:.0%}, {var_result.horizon_days}d"),
        ("CVaR", f"{var_result.cvar * 100:.2f}%", ""),
        ("Portfolio delta", f"{greeks_result.portfolio_delta:,.0f}", ""),
        ("ESG score", f"{sustainability_result.portfolio_esg_score:.1f}", "0–100 weighted average"),
        ("Carbon intensity", f"{sustainability_result.carbon_intensity_weighted:.1f}", "tCO2e/$M"),
        ("Carbon VaR", f"{sustainability_result.carbon_var * 100:.2f}%", "of portfolio MV"),
        ("Stress P&L", f"{stress_result.portfolio_pct * 100:.2f}%", stress_result.scenario),
    ]
    ws.cell(row=2, column=1, value="Metric").font = _BOLD_HEADER
    ws.cell(row=2, column=2, value="Value").font = _BOLD_HEADER
    ws.cell(row=2, column=3, value="Notes").font = _BOLD_HEADER
    for label, value, note in metrics:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=3, value=note)
        row += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 36

    position_rows = []
    for pos in portfolio.positions:
        mv = pos.current_price * pos.quantity if pos.asset_class.value != "Cash" else pos.quantity
        position_rows.append(
            {
                "ticker": pos.ticker,
                "asset_class": pos.asset_class.value,
                "quantity": pos.quantity,
                "current_price": pos.current_price,
                "market_value": mv,
                "esg_score": pos.esg_score,
                "beta": pos.beta,
            }
        )
    if position_rows:
        _write_detail_sheet(wb, position_rows, title="Holdings")

    var_detail = [
        {
            "ticker": c.ticker,
            "var_contribution": c.var_contribution,
            "pct_of_total": c.pct_of_total * 100.0,
        }
        for c in var_result.contributions
    ]
    if var_detail:
        _write_detail_sheet(wb, var_detail, title="VaR Detail")

    _write_model_parameters_sheet(wb, model_parameters, path, extra={"report_type": "full_risk"})
    if chart_images:
        _write_charts_sheet(wb, chart_images)

    wb.save(path)
    return str(path.resolve())


def validate_report_path(path: str) -> str:
    """Reject paths outside engine/output/reports (defence in depth for desktop open)."""
    resolved = Path(path).resolve()
    reports_root = _REPORTS_DIR.resolve()
    reports_root.mkdir(parents=True, exist_ok=True)
    try:
        resolved.relative_to(reports_root)
    except ValueError as exc:
        raise ValueError("Report path is outside the allowed reports directory.") from exc
    if resolved.suffix.lower() != ".xlsx":
        raise ValueError("Report path must be an .xlsx file.")
    if not resolved.is_file():
        raise ValueError("Report file was not found.")
    return str(resolved)


def _write_detail_sheet(
    wb: Workbook,
    detail_rows: list[dict[str, Any]],
    *,
    title: str = "Position Detail",
) -> None:
    if not detail_rows:
        return
    detail = wb.create_sheet(title)
    headers = list(detail_rows[0].keys())
    for col, h in enumerate(headers, start=1):
        detail.cell(row=1, column=col, value=excel_safe_cell(h)).font = _BOLD_HEADER
    for r_idx, row_data in enumerate(detail_rows[:_MAX_DETAIL_ROWS], start=2):
        for c_idx, h in enumerate(headers, start=1):
            val = row_data.get(h)
            detail.cell(
                row=r_idx,
                column=c_idx,
                value=excel_safe_cell(val) if isinstance(val, str) else val,
            )


def _write_model_parameters_sheet(
    wb: Workbook,
    model_parameters: dict[str, Any],
    path: Path,
    *,
    extra: dict[str, str] | None = None,
) -> None:
    params = wb.create_sheet("Model Parameters")
    params["A1"] = "Generated (UTC)"
    params["B1"] = datetime.now(timezone.utc).isoformat()
    params["A2"] = "Report path"
    params["B2"] = str(path.resolve())
    row = 3
    for key, value in (extra or {}).items():
        params.cell(row=row, column=1, value=excel_safe_cell(key))
        params.cell(row=row, column=2, value=excel_safe_cell(value))
        row += 1
    for key, value in sorted(model_parameters.items()):
        params.cell(row=row, column=1, value=excel_safe_cell(key))
        params.cell(row=row, column=2, value=excel_safe_cell(value))
        row += 1
    params.column_dimensions["A"].width = 28
    params.column_dimensions["B"].width = 60


def _write_charts_sheet(wb: Workbook, chart_images: list[tuple[str, bytes]]) -> None:
    ws = wb.create_sheet("Charts")
    row = 1
    for title, png in chart_images[:_MAX_CHARTS]:
        if not _is_valid_png(png):
            continue
        ws.cell(row=row, column=1, value=excel_safe_cell(title)).font = _BOLD_HEADER
        row += 1
        img = Image(io.BytesIO(png))
        img.width = 640
        img.height = 360
        ws.add_image(img, f"A{row}")
        row += 22
    ws.column_dimensions["A"].width = 80


def _slugify(name: str) -> str:
    import re

    slug = re.sub(r"[^a-z0-9]+", "_", name.lower())[:30].strip("_")
    return slug or "portfolio"


def _is_valid_png(png: bytes) -> bool:
    if (
        len(png) < len(_PNG_MAGIC)
        or not png.startswith(_PNG_MAGIC)
        or len(png) > _MAX_CHART_PNG_BYTES
    ):
        return False
    try:
        from PIL import Image

        with Image.open(io.BytesIO(png)) as img:
            img.verify()
        return True
    except Exception:
        return False

